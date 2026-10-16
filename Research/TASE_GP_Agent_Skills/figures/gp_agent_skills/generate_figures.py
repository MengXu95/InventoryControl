"""Native editable GP / Agent / Skill diagrams; export with local PowerPoint COM."""

from __future__ import annotations

import argparse
import json
import math
import re
import subprocess
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

import pymupdf
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
DECK = ROOT / "gp_agent_skills_figures.pptx"
WIDTH, HEIGHT = 16.0, 9.0
PAPER_WIDTH = 7.125
PAPER_TITLE = ("PACT-GP: Genetic Programming for Skill-Based Agent Coordination "
               "in Multi-Site Inventory Control")
FONT = "Arial"
INK = "203047"
MUTED = "536477"
WHITE = "FFFFFF"
OFFLINE = "C87524"
OFFLINE_BG = "FFF3E7"
AGENT = "2369A4"
AGENT_BG = "EDF5FC"
PROTOCOL = "287B5A"
PROTOCOL_BG = "EDF8F1"
ACTUAL = "8054A7"
ACTUAL_BG = "F4EEF9"
BORDER = "D8E0E8"
LIGHT = "F7F9FC"


@dataclass
class Shape:
    kind: str
    x: float
    y: float
    w: float
    h: float
    text: str = ""
    font: float = 17
    fill: str | None = None
    stroke: str | None = None
    color: str = INK
    bold: bool = False
    align: str = "center"
    padding: float = 0.12
    arrow: bool = False
    reverse_arrow: bool = False
    thickness: float = 1.5
    semantic: str = ""


@dataclass
class Figure:
    name: str
    title: str
    subtitle: str
    expected: list[str]
    shapes: list[Shape] = field(default_factory=list)

    def text(self, x, y, w, h, text, font=17, color=INK, bold=False,
             align="left", padding=0.0, semantic=""):
        self.shapes.append(Shape("text", x, y, w, h, text, font,
                                 color=color, bold=bold, align=align,
                                 padding=padding, semantic=semantic))

    def box(self, x, y, w, h, text="", font=17, fill=WHITE, stroke=BORDER,
            color=INK, bold=False, align="center", padding=0.12,
            kind="round", semantic=""):
        self.shapes.append(Shape(kind, x, y, w, h, text, font, fill,
                                 stroke, color, bold, align, padding,
                                 semantic=semantic))

    def line(self, x1, y1, x2, y2, color=MUTED, arrow=True, both=False,
             thickness=1.7):
        self.shapes.append(Shape("line", x1, y1, x2-x1, y2-y1,
                                 stroke=color, arrow=arrow,
                                 reverse_arrow=both, thickness=thickness))

    def path(self, points, color=MUTED, arrow=True, thickness=1.7):
        for i, (a, b) in enumerate(zip(points, points[1:])):
            self.line(*a, *b, color=color,
                      arrow=arrow and i == len(points)-2, thickness=thickness)

    def base(self, index):
        self.box(0, 0, WIDTH, HEIGHT, fill=WHITE, stroke=None,
                 kind="rect", padding=0)
        self.text(.5, .32, 14.6, .53, self.title, 29, bold=True)
        self.text(.5, .96, 15, .56, self.subtitle, 18, MUTED)
        self.line(.5, 1.58, 15.5, 1.58, BORDER, False, thickness=1)


def architecture():
    f = Figure("architecture", "PACT-GP: one shared policy, many local peer agents",
               "Offline evolution changes expression trees. Online Skills apply the frozen "
               "trees to different private local inputs.",
               ["OFFLINE", "ONLINE", "Same frozen individual",
                "Agent A", "Agent B", "Agent C", "No LLM"])
    f.base(1)
    f.box(.5, 1.88, 3.25, 5.87, fill=OFFLINE_BG, stroke=OFFLINE)
    f.text(.72, 2.07, 2.83, .38, "OFFLINE  |  GP search", 18, OFFLINE, True)
    f.box(.76, 2.77, 2.73, 1.20,
          "One GP individual g\n(T_order, T_request, T_offer)",
          17, WHITE, OFFLINE, bold=True)
    f.box(.76, 4.58, 2.73, 1.25,
          "Centralized simulator\nSeeded episodes\nFitness / genetic variation",
          16, WHITE, OFFLINE)
    f.line(2.12, 4.58, 2.12, 3.97, OFFLINE)
    f.text(.80, 6.21, 2.65, 1.02,
           "Evolve three trees.\nFreeze validated g.\nNo online controller.",
           15, MUTED, align="center")
    f.box(4.18, 1.88, 11.32, 5.87, fill=WHITE, stroke=AGENT)
    f.text(4.45, 2.07, 10.7, .38,
           "ONLINE  |  decentralized deterministic Skills", 18, AGENT, True)
    f.box(4.45, 2.65, 10.75, .55,
          "Same frozen individual at every site; local inputs and state differ",
          17, OFFLINE_BG, OFFLINE, color=OFFLINE, bold=True)
    f.line(3.49, 3.37, 4.45, 2.94, OFFLINE)
    for x, label in [(4.48, "A"), (8.12, "B"), (11.76, "C")]:
        f.box(x, 3.63, 3.13, 3.55, fill=AGENT_BG, stroke=AGENT)
        f.text(x+.15, 3.79, 2.83, .36, f"Agent {label}", 20, AGENT, True,
               align="center")
        f.text(x+.16, 4.26, 2.81, .54,
               "Own inventory, pipeline,\nforecast and costs", 15, align="center")
        f.box(x+.16, 4.98, 2.81, .51,
              "Shared three-tree g",
              12.5, OFFLINE_BG, OFFLINE, color=OFFLINE, bold=True, padding=.06)
        f.line(x+1.57, 3.20, x+1.57, 3.63, OFFLINE)
        f.box(x+.16, 5.73, 2.81, .78, "Skills\norder / request / offer",
              13.5, WHITE, AGENT, color=AGENT, bold=True, padding=.06)
        f.line(x+1.57, 5.49, x+1.57, 5.73, OFFLINE)
        f.text(x+.16, 6.58, 2.81, .55, "Reservations\nActual-flow history",
               12.5, align="center")
    f.line(7.63, 6.00, 8.10, 6.00, PROTOCOL, both=True)
    f.line(11.27, 6.00, 11.74, 6.00, PROTOCOL, both=True)
    f.text(4.56, 7.35, 10.4, .36,
           "Peer request / offer / confirm messages — no master / subordinate relation",
           16, PROTOCOL, True, align="center")
    f.box(.5, 8.02, 15, .37,
          "Fixed, not evolved: role selection • bounds • gain proxy • mutual-best rounds "
          "• reserve / confirm • physical execution.   No LLM.",
          13, LIGHT, None, color=MUTED, padding=.04)
    return f


def negotiation():
    f = Figure("negotiation", "Peer matching: two receivers compete for one donor",
               "Illustrative K = 0 example (simulator default K = 10): G(q) = 1.5q. "
               "All promises are pre-demand.",
               ["ROUND 1", "ROUND 2", "Donor D", "Receiver A", "Receiver B",
                "D–B: 8", "D–A: 2", "not dispatch", "simulator default K = 10",
                "G(q) = 1.5q"])
    f.base(2)
    f.box(.5, 1.88, 10.2, 3.89, fill=LIGHT, stroke=BORDER)
    f.text(.76, 2.05, 9.65, .32, "ROUND 1  |  A and B both request from D",
           18, PROTOCOL, True)
    f.box(4.16, 2.68, 2.85, 1.02, "Donor D\nOffer capacity = 10",
          18, AGENT_BG, AGENT, bold=True)
    f.box(.86, 4.40, 2.85, 1.02, "Receiver A\nRequest = 6",
          18, AGENT_BG, AGENT, bold=True)
    f.box(7.50, 4.40, 2.85, 1.02, "Receiver B\nRequest = 8",
          18, AGENT_BG, AGENT, bold=True)
    f.line(3.56, 4.40, 4.50, 3.70, PROTOCOL, both=True, thickness=2)
    f.line(6.70, 3.70, 7.66, 4.40, PROTOCOL, both=True, thickness=2.6)
    f.text(.90, 3.30, 2.95, .81, "q = 6; gain = 9\nA chooses D\nD rejects A",
           13.5, PROTOCOL, align="center")
    f.text(7.43, 3.30, 2.95, .81, "q = 8; gain = 12\nD and B choose\none another",
           13.5, PROTOCOL, True, align="center")
    f.box(4.16, 4.52, 2.85, .78, "Mutual best: D–B\nReserve 8 each",
          16, PROTOCOL_BG, PROTOCOL, color=PROTOCOL, bold=True)
    f.line(5.58, 3.70, 5.58, 4.52, PROTOCOL)
    f.box(.5, 6.02, 10.2, 2.31, fill=LIGHT, stroke=BORDER)
    f.text(.76, 6.18, 9.65, .32,
           "ROUND 2  |  reservations leave D with 2 units", 18, PROTOCOL, True)
    f.box(.86, 6.82, 2.85, .90, "Receiver A\nRequest left = 6",
          16, AGENT_BG, AGENT, bold=True)
    f.box(4.16, 6.82, 2.85, .90, "Donor D\nOffer left = 2",
          16, AGENT_BG, AGENT, bold=True)
    f.box(7.50, 6.82, 2.85, .90, "Receiver B\nRequest left = 0",
          16, LIGHT, BORDER, color=MUTED)
    f.line(3.73, 7.26, 4.14, 7.26, PROTOCOL, both=True)
    f.text(.87, 7.91, 9.4, .24,
           "D–A: q = 2, gain = 3; mutual best → reserve 2.  B has no residual request.",
           13.5, PROTOCOL, align="center")
    f.box(11.03, 1.88, 4.47, 6.45, fill=PROTOCOL_BG, stroke=PROTOCOL)
    f.text(11.29, 2.10, 3.95, .38, "Fixed matching rules", 21, PROTOCOL, True)
    f.text(11.29, 2.76, 3.94, 1.59,
           "• Common total edge ranking\n• Choose positive-gain best edge\n• Confirm only mutual choices\n• At most one partner per round\n• Reserve before the next round",
           17)
    f.line(11.29, 4.63, 15.24, 4.63, PROTOCOL, False, thickness=1)
    f.text(11.29, 4.87, 3.95, .35, "Conditional promise ledger", 18,
           PROTOCOL, True)
    f.box(11.30, 5.48, 3.93, 1.02,
          "D–B: 8     D–A: 2\nTotal reserved at D = 10",
          18, WHITE, PROTOCOL, bold=True)
    f.text(11.29, 6.78, 3.95, 1.14,
           "Promises are not dispatch.\nNo receive-forward.\nRecheck after demand.",
           16, color=MUTED)
    return f


def gp_search():
    f = Figure("gp_search", "PACT-GP: offline search and leakage-free policy selection",
               "The genotype is a three-tree individual. The entire fixed online protocol "
               "is executed inside every seeded training episode.",
               ["Initialize population", "Common seeded episodes", "Site service constraints",
                "Same-role", "Validation select", "Frozen test", "No LLM"])
    f.base(3)
    x, w = .60, 4.20
    f.box(x, 1.96, w, .88, "Initialize population\ng = (T_order, T_request, T_offer)",
          17, OFFLINE_BG, OFFLINE, bold=True)
    f.box(x, 3.21, w, .99, "Common seeded episodes\nShared scenario / demand seeds\nAll local agents + fixed protocol",
          17, OFFLINE_BG, OFFLINE, padding=.10)
    f.line(2.70, 2.84, 2.70, 3.21, OFFLINE)
    f.box(x, 4.57, w, 1.04,
          "Site service constraints first\nCost ranks admissible policies",
          17, OFFLINE_BG, OFFLINE, bold=True)
    f.line(2.70, 4.20, 2.70, 4.57, OFFLINE)
    f.box(x, 6.29, w, 1.03, "Stop criterion met?\nGeneration / evaluation budget",
          18, WHITE, OFFLINE, bold=True)
    f.line(2.70, 5.61, 2.70, 6.29, OFFLINE)
    f.box(6.30, 4.63, 3.55, .95, "Selection\nChoose parent individuals",
          18, OFFLINE_BG, OFFLINE, bold=True)
    f.box(11.05, 4.63, 4.35, .95,
          "Same-role crossover / mutation\nOrder × order; request × request;\noffer × offer",
          16, OFFLINE_BG, OFFLINE, bold=True, padding=.06)
    f.line(9.85, 5.10, 11.05, 5.10, OFFLINE)
    f.box(11.05, 2.73, 4.35, 1.03,
          "Enforce structural limits\nTree depth / size; valid primitives\nRepair or reject invalid offspring",
          16, OFFLINE_BG, OFFLINE, bold=True)
    f.line(13.23, 4.63, 13.23, 3.76, OFFLINE)
    f.path([(11.05, 3.25), (10.38, 3.25), (10.38, 3.70), (4.80, 3.70)],
           OFFLINE)
    f.text(6.01, 3.86, 3.90, .30, "Reevaluate on common seeds",
           14, OFFLINE, align="center")
    f.path([(3.85, 6.29), (3.85, 5.94), (8.08, 5.94), (8.08, 5.58)],
           OFFLINE)
    f.text(4.91, 5.99, 1.12, .26, "No", 13, OFFLINE)
    f.box(6.30, 6.29, 3.55, 1.03,
          "Validation select\nSelect g on held-out data\nFreeze all three trees",
          17, AGENT_BG, AGENT, bold=True)
    f.line(4.80, 6.80, 6.30, 6.80, AGENT)
    f.text(5.04, 6.39, .93, .25, "Yes", 14, AGENT, True, align="center")
    f.box(11.05, 6.29, 4.35, 1.03,
          "Frozen test → deploy\nReport unseen test performance\nNo tuning on test episodes",
          17, AGENT_BG, AGENT, bold=True)
    f.line(9.85, 6.80, 11.05, 6.80, AGENT)
    f.text(6.16, 1.99, 4.18, .88,
           "Study budget: P >= 200; G >= 50.\nDisjoint train / validation / test.\nNo LLM in either phase.",
           16, MUTED)
    f.box(.60, 7.76, 14.80, .49,
          "Only the three GP trees evolve; roles, bounds, gain evaluation, matching, "
          "reservations and physical execution stay fixed.",
          14, LIGHT, None, color=MUTED)
    return f


def tree(f, cx, role, offer=False):
    # Node positions encode left-associative subtraction exactly as shown.
    if not offer:
        nodes = [
            ("max", cx, 3.44), ("0", cx-1.30, 4.32),
            ("−", cx+.65, 4.32), ("−", cx+.05, 5.20),
            ("P", cx+1.55, 5.20), ("+", cx-.65, 6.08),
            ("I", cx+.78, 6.08), ("F1", cx-1.35, 6.96),
            ("F2", cx+.03, 6.96),
        ]
        edges = [(0, 1), (0, 2), (2, 3), (2, 4),
                 (3, 5), (3, 6), (5, 7), (5, 8)]
    else:
        nodes = [
            ("max", cx, 3.44), ("0", cx-1.30, 4.32),
            ("−", cx+.65, 4.32), ("−", cx+.05, 5.20),
            ("F2", cx+1.55, 5.20), ("I", cx-.65, 6.08),
            ("F1", cx+.78, 6.08),
        ]
        edges = [(0, 1), (0, 2), (2, 3), (2, 4), (3, 5), (3, 6)]
    radii = [.42 if label == "max" else .285 for label, _, _ in nodes]
    for a, b in edges:
        _, ax, ay = nodes[a]
        _, bx, by = nodes[b]
        length = math.hypot(bx-ax, by-ay)
        dx, dy = (bx-ax)/length, (by-ay)/length
        f.line(ax+dx*radii[a], ay+dy*radii[a],
               bx-dx*radii[b], by-dy*radii[b], OFFLINE, arrow=False)
    for (label, nx, ny), radius in zip(nodes, radii):
        f.box(nx-radius, ny-radius, radius*2, radius*2, label, 16,
              WHITE, OFFLINE, color=OFFLINE, bold=True, padding=.015,
              kind="ellipse", semantic=f"tree:{role}:{label}")


def gp_individual():
    f = Figure("gp_individual", "A GP individual contains three role-specific trees",
               "Illustrative valid expressions, not initialized seed or trained optimum. "
               "All nodes and edges are editable.",
               ["T_order", "T_request", "T_offer", "Illustrative valid expressions",
                "max(0, F1 + F2 − I − P)", "max(0, I − F1 − F2)"])
    f.base(4)
    columns = [
        (.50, "T_order", "max(0, F1 + F2 − I − P)", False,
         "Fixed order / capacity caps"),
        (5.65, "T_request", "max(0, F1 + F2 − I − P)", False,
         "Fixed receiver / request cap"),
        (10.80, "T_offer", "max(0, I − F1 − F2)", True,
         "Fixed donor / offer cap"),
    ]
    for x, role, formula, offer, cap in columns:
        f.box(x, 1.91, 4.70, 6.42, fill=OFFLINE_BG, stroke=OFFLINE)
        f.text(x+.16, 2.08, 4.38, .41, role, 23, OFFLINE, True, align="center")
        f.text(x+.10, 2.61, 4.50, .31, formula, 17, INK, True, align="center")
        tree(f, x+2.12, role, offer)
        f.box(x+.16, 7.66, 4.38, .49, cap, 13, WHITE, AGENT, color=AGENT,
              bold=True, padding=.04)
    f.text(.55, 8.36, 14.8, .21,
           "I: on-hand inventory   P: own pipeline total   F1, F2: exposed local forecasts.  "
           "The fixed layer selects the active donor / receiver role.",
           11, MUTED, align="center")
    return f


def timing():
    f = Figure("timing", "Conditional promises are not realized transshipments",
               "Negotiation precedes current demand. Physical dispatch is decided only by "
               "the fixed post-demand feasibility recheck.",
               ["PRE-DEMAND", "due arrivals", "current demand", "Full", "Partial",
                "Cancelled", "actual flow only", "No post-demand oracle"])
    f.base(5)
    f.text(.55, 1.81, 14.8, .35,
           "PRE-DEMAND  |  local information + exposed forecasts + received peer messages",
           18, PROTOCOL, True)
    top = [
        (.55, "1  Local observation", "Own inventory / pipeline\nForecasts / local costs"),
        (4.32, "2  Deterministic Skills", "Evaluate shared frozen GP\nApply fixed roles and caps"),
        (8.09, "3  R peer rounds", "Positive-gain mutual best\nReserve and confirm"),
        (11.86, "4  Conditional promises", "Commit q_commit by partner\nNo physical dispatch yet"),
    ]
    for x, title, detail in top:
        f.box(x, 2.47, 3.59, 1.24, fill=PROTOCOL_BG, stroke=PROTOCOL)
        f.text(x+.12, 2.64, 3.35, .32, title, 17, PROTOCOL, True, align="center")
        f.text(x+.12, 3.04, 3.35, .55, detail, 15, align="center")
    for x1, x2 in [(4.14, 4.32), (7.91, 8.09), (11.68, 11.86)]:
        f.line(x1, 3.10, x2, 3.10, PROTOCOL)
    f.box(.55, 4.03, 14.90, .59,
          "EVENT BOUNDARY: due arrivals → current demand → actual inventory and unmet demand",
          18, ACTUAL_BG, ACTUAL, color=ACTUAL, bold=True)
    f.text(.55, 4.85, 6.4, .31, "POST-DEMAND  |  fixed physical execution", 18,
           ACTUAL, True)
    f.box(.55, 5.55, 3.35, 1.34,
          "5  Feasibility recheck\nActual donor / receiver state\n0 ≤ q_exec ≤ q_commit",
          15, ACTUAL_BG, ACTUAL, bold=True)
    f.box(4.52, 5.55, 2.32, 1.34,
          "6  Realized flow\nq_exec from actual\nfeasible bounds",
          16, ACTUAL_BG, ACTUAL, bold=True)
    f.line(3.90, 6.22, 4.52, 6.22, ACTUAL)
    outcomes = [
        (5.17, "Full", "q_exec = q_commit > 0"),
        (6.27, "Partial", "0 < q_exec < q_commit"),
        (7.37, "Cancelled", "q_exec = 0"),
    ]
    for y, title, expr in outcomes:
        f.box(7.62, y, 3.18, .83, f"{title}\n{expr}", 16,
              ACTUAL_BG, ACTUAL, bold=True)
        f.path([(6.84, 6.22), (7.20, 6.22), (7.20, y+.415), (7.62, y+.415)],
               ACTUAL, thickness=1.3)
        f.path([(10.80, y+.415), (11.15, y+.415), (11.15, 6.63), (11.48, 6.63)],
               ACTUAL, thickness=1.3)
    f.box(11.48, 5.46, 3.97, 2.35,
          "7  State and accounting\nCosts + losses\nPipeline update\nRelease reservations\nHistory: actual flow only",
          18, ACTUAL_BG, ACTUAL, bold=True)
    f.text(.65, 7.36, 6.20, .78,
           "No post-demand oracle during negotiation.\nA promise can execute fully, partly or not at all.",
           17, MUTED)
    return f


def fixed_evaluator():
    f = Figure("fixed_evaluator", "Fixed local gain evaluator and peer message interface",
               "Conservative forecast caps keep Φ on linear branches. "
               "Heuristic ranking, not an expected-cost guarantee.",
               ["ρ = I + Σ pipeline − Σ exposed forecasts", "Φ(ρ)",
                "A(q)", "B(q)", "G(q)", "not an expected-cost guarantee",
                "reservation token", "No receive-forward", "no interior optimum",
                "(h_donor + b_receiver − c) q − K"])
    f.base(6)
    f.box(.5, 1.90, 8.60, 6.43, fill=LIGHT, stroke=BORDER)
    f.text(.76, 2.07, 8.08, .34, "Fixed forecast-slack proxy", 21, PROTOCOL, True)
    f.box(.77, 2.66, 8.06, .93,
          "ρ = I + Σ pipeline − Σ exposed forecasts\nρ > 0: donor;  ρ < 0: receiver;  ρ = 0: inactive",
          19, WHITE, PROTOCOL, bold=True)
    f.box(.77, 3.82, 8.06, .65,
          "Φ(ρ) = h max(ρ, 0) + b max(−ρ, 0)",
          21, WHITE, PROTOCOL, bold=True)
    f.box(.77, 4.73, 3.91, .95,
          "Donor marginal cost\nA(q) = Φ(ρ_d − q) − Φ(ρ_d)",
          16, WHITE, PROTOCOL, bold=True)
    f.box(4.92, 4.73, 3.91, .95,
          "Receiver marginal benefit\nB(q) = Φ(ρ_r) − Φ(ρ_r + q)",
          16, WHITE, PROTOCOL, bold=True)
    f.box(.77, 5.94, 8.06, .92,
          "G(q) = B(q) − A(q) − C_trans(q)\n"
          "= (h_donor + b_receiver − c) q − K,  q > 0",
          19, PROTOCOL_BG, PROTOCOL, color=PROTOCOL, bold=True)
    f.text(.85, 7.02, 7.91, .59,
           "Caps: request ≤ forecast deficit; offer ≤ forecast surplus.\n"
           "A positive-gain pair takes largest q; no interior optimum.", 18,
           MUTED, align="center")
    f.box(.77, 7.78, 8.06, .51,
          "q ∈ {0, 0.25, 0.5, 0.75, 1} × min(remaining offer, remaining request)",
          15, WHITE, PROTOCOL, padding=.05)
    f.box(9.50, 1.90, 6.00, 3.85, fill=AGENT_BG, stroke=AGENT)
    f.text(9.77, 2.08, 5.46, .36, "Local message schema (illustrative)", 21,
           AGENT, True)
    f.text(9.77, 2.78, 5.45, 2.15,
           "Header: site / peer IDs, period, round\nRole: request or offer\nRemaining request / offer quantity\nCandidate q and local marginal values\nCommon ranking / deterministic tie key\nConfirm or release; reservation token",
           18)
    f.text(9.77, 5.12, 5.45, .31,
           "Private inventory, pipeline, forecasts.", 15, AGENT, True)
    f.box(9.50, 6.03, 6.00, 2.30, fill=PROTOCOL_BG, stroke=PROTOCOL)
    f.text(9.77, 6.20, 5.46, .34, "Fixed protocol invariants", 21, PROTOCOL, True)
    f.text(9.77, 6.83, 5.45, 1.18,
           "• Common total ranking of positive-gain edges\n• Mutual best; at most one partner / round\n• Reserve / confirm prevents overpromise\n• No receive-forward within the period",
           17)
    return f


def figures():
    items = [architecture(), negotiation(), gp_search(),
             gp_individual(), timing(), fixed_evaluator()]
    for figure in items:
        for shape in figure.shapes:
            if shape.text:
                shape.font = max(shape.font, 18)
                shape.padding = min(shape.padding, .08)
                if shape.h < .36:
                    shape.y -= (.36-shape.h)/2
                    shape.h = .36
                if shape.kind == "text":
                    needed = len(paragraph_lines(shape))*shape.font*1.30/72+.03
                    if needed > shape.h:
                        shape.y -= (needed-shape.h)/2
                        shape.h = needed
    return items


def rgb(value):
    return RGBColor.from_string(value)


def paragraph_lines(shape):
    font_file = Path(r"C:\Windows\Fonts") / ("arialbd.ttf" if shape.bold else "arial.ttf")
    metric = ImageFont.truetype(str(font_file), round(shape.font*4))
    available = (shape.w-2*shape.padding)*72*4
    all_lines = []
    for explicit in shape.text.splitlines():
        current = ""
        for word in explicit.split(" "):
            proposed = f"{current} {word}" if current else word
            if metric.getlength(proposed) <= available:
                current = proposed
            elif not current:
                raise AssertionError(f"Word exceeds textbox width: {shape.text!r}")
            else:
                all_lines.append(current)
                current = word
        all_lines.append(current)
    return all_lines


def validate_model(items):
    fit_errors = []
    for figure in items:
        for shape in figure.shapes:
            if shape.kind == "line":
                xs = [shape.x, shape.x+shape.w]
                ys = [shape.y, shape.y+shape.h]
            else:
                xs = [shape.x, shape.x+shape.w]
                ys = [shape.y, shape.y+shape.h]
            assert min(xs) >= -.001 and max(xs) <= WIDTH+.001, (
                figure.name, "horizontal bounds", shape)
            assert min(ys) >= -.001 and max(ys) <= HEIGHT+.001, (
                figure.name, "vertical bounds", shape)
            if shape.text:
                lines = paragraph_lines(shape)
                required = len(lines)*shape.font*1.08
                available = (shape.h-2*shape.padding)*72
                if required > available+.5:
                    fit_errors.append(
                        f"{figure.name}: text needs {required:.1f}pt, has {available:.1f}pt: "
                        f"{shape.text!r}; wrapped={lines}")
        text = "\n".join(s.text for s in figure.shapes)
        for expected in figure.expected:
            assert expected in text, (figure.name, expected)
    counts = {role: sum(s.semantic.startswith(f"tree:{role}:")
                       for f in items for s in f.shapes)
              for role in ("T_order", "T_request", "T_offer")}
    assert counts == {"T_order": 9, "T_request": 9, "T_offer": 7}, counts
    assert not fit_errors, "\n" + "\n".join(fit_errors)


def set_arrow(line, head, tail):
    ln = line._get_or_add_ln()
    for name, enabled in (("headEnd", tail), ("tailEnd", head)):
        element = OxmlElement(f"a:{name}")
        element.set("type", "triangle" if enabled else "none")
        element.set("w", "sm")
        element.set("len", "sm")
        ln.append(element)


def write_deck(items):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(WIDTH), Inches(HEIGHT)
    deck.core_properties.title = PAPER_TITLE
    deck.core_properties.subject = "PACT-GP: peer-agent coordination with a shared three-tree policy"
    deck.core_properties.author = "InventoryControl research"
    deck.core_properties.keywords = "PACT-GP, genetic programming, agent, skill, peer negotiation"
    for figure in items:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.name = figure.name
        for i, item in enumerate(figure.shapes, 1):
            coords = [Inches(v) for v in (item.x, item.y, item.w, item.h)]
            if item.kind == "line":
                shape = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT, coords[0], coords[1],
                    Inches(item.x+item.w), Inches(item.y+item.h))
                shape.line.color.rgb = rgb(item.stroke or MUTED)
                shape.line.width = Pt(item.thickness)
                set_arrow(shape.line, item.arrow, item.reverse_arrow)
            else:
                if item.kind == "text":
                    shape = slide.shapes.add_textbox(*coords)
                else:
                    kind = {"rect": MSO_SHAPE.RECTANGLE,
                            "round": MSO_SHAPE.ROUNDED_RECTANGLE,
                            "ellipse": MSO_SHAPE.OVAL}[item.kind]
                    shape = slide.shapes.add_shape(kind, *coords)
                    if item.kind == "round":
                        shape.adjustments[0] = .09
                    if item.fill:
                        shape.fill.solid()
                        shape.fill.fore_color.rgb = rgb(item.fill)
                    else:
                        shape.fill.background()
                    if item.stroke:
                        shape.line.color.rgb = rgb(item.stroke)
                        shape.line.width = Pt(1.0)
                    else:
                        shape.line.fill.background()
                tf = shape.text_frame
                tf.clear()
                tf.word_wrap = item.kind != "ellipse"
                tf.margin_left = tf.margin_right = Inches(item.padding)
                tf.margin_top = tf.margin_bottom = Inches(item.padding)
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                for index, text in enumerate(item.text.split("\n")):
                    p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
                    p.text = text
                    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                                   "right": PP_ALIGN.RIGHT}[item.align]
                    p.space_before = p.space_after = Pt(0)
                    p.line_spacing = 1.08
                    for run in p.runs:
                        run.font.name = FONT
                        run.font.size = Pt(item.font)
                        run.font.bold = item.bold
                        run.font.color.rgb = rgb(item.color)
            shape._element.spPr.append(OxmlElement("a:effectLst"))
            shape.name = f"{figure.name}:{i:03d}:{item.semantic or item.kind}"
        slide.notes_slide.notes_text_frame.text = (
            figure.title + "\n" + figure.subtitle + "\n"
            "Native editable PowerPoint geometry and text. No LLM in either phase. "
            "Orange: offline GP; blue: site agents / Skills; green: fixed protocol; "
            "purple: actual execution. See README.md for notation and regeneration.")
    deck.save(DECK)


NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def validate_deck(items):
    summary = []
    with zipfile.ZipFile(DECK) as archive:
        slides = sorted((n for n in archive.namelist()
                         if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                        key=lambda n: int(re.search(r"slide(\d+)", n).group(1)))
        assert len(slides) == len(items) == 6
        media = [n for n in archive.namelist() if n.startswith("ppt/media/")]
        assert media == [], f"No images expected, found {media}"
        for item, path in zip(items, slides):
            xml = ET.fromstring(archive.read(path))
            text = "\n".join(el.text or "" for el in xml.findall(".//a:t", NS))
            for expected in item.expected:
                assert expected in text, (item.name, expected)
            count = len(xml.findall(".//p:sp", NS))
            connectors = len(xml.findall(".//p:cxnSp", NS))
            assert count+connectors == len(item.shapes)
            assert not xml.findall(".//p:pic", NS)
            for node in xml.findall(".//p:sp", NS):
                name = node.find("p:nvSpPr/p:cNvPr", NS).get("name", "")
                if ":tree:" in name:
                    assert node.find("p:txBody/a:bodyPr", NS).get("wrap") == "none", (
                        item.name, "Tree labels must not wrap", name)
            summary.append({"figure": item.name, "native_shapes": count,
                            "native_connectors": connectors,
                            "native_text_runs": len(xml.findall(".//a:t", NS)),
                            "pictures": 0})
    return summary


def normalized(text):
    return re.sub(r"\s+", "", text).replace("−", "-").replace("–", "-").casefold()


def validate_pdfs(items, preview):
    results = []
    fit_errors = []
    for item in items:
        path = ROOT / f"{item.name}.pdf"
        assert path.exists(), f"Missing vector export: {path}"
        with pymupdf.open(path) as doc:
            assert len(doc) == 1, path
            assert doc.metadata["title"] == PAPER_TITLE, (path.name, doc.metadata["title"])
            page = doc[0]
            if preview:
                page.get_pixmap(matrix=pymupdf.Matrix(1.25, 1.25),
                                alpha=False).save(str(ROOT / f"{item.name}_preview.png"))
            assert abs(page.rect.width-WIDTH*72) < 1
            assert abs(page.rect.height-HEIGHT*72) < 1
            assert not page.get_images(), f"Raster images found in {path}"
            vector_count = len(page.get_drawings())
            assert vector_count >= 10, (path, vector_count)
            text = page.get_text()
            for expected in item.expected:
                assert normalized(expected) in normalized(text), (
                    path.name, "Missing rendered text", expected)
            for shape in item.shapes:
                if shape.text:
                    assert normalized(shape.text) in normalized(text), (
                        path.name, "Incomplete rendered textbox", shape.text)
            spans = []
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        bounds = pymupdf.Rect(span["bbox"])
                        assert page.rect.contains(bounds), (path.name, span)
                        spans.append(span)
            smallest_font = min(span["size"] for span in spans if span["text"].strip())
            article_font = smallest_font*PAPER_WIDTH/WIDTH
            assert smallest_font >= 17.99 and article_font >= 8.0, (
                path.name, "Article-scale text too small", smallest_font, article_font)
            for shape in item.shapes:
                if shape.semantic.startswith("tree:"):
                    node_bounds = pymupdf.Rect(shape.x*72, shape.y*72,
                                               (shape.x+shape.w)*72, (shape.y+shape.h)*72)
                    labels = [
                        span for span in spans
                        if span["text"].strip() == shape.text
                        and node_bounds.contains(pymupdf.Rect(span["bbox"]))
                    ]
                    assert len(labels) == 1, (
                        path.name, "Tree label must render as one complete span",
                        shape.semantic, labels)
            # Check visible text against its own native textbox, and different textboxes.
            owners = []
            for span in spans:
                bbox = pymupdf.Rect(span["bbox"])
                center = pymupdf.Point((bbox.x0+bbox.x1)/2, (bbox.y0+bbox.y1)/2)
                candidates = [
                    (i, s) for i, s in enumerate(item.shapes) if s.text
                    and pymupdf.Rect(s.x*72, s.y*72, (s.x+s.w)*72, (s.y+s.h)*72)
                    .contains(center)
                    and normalized(span["text"]) in normalized(s.text)
                ]
                assert candidates, (path.name, "Text outside expected textbox", span["text"])
                owner, shape = min(candidates, key=lambda pair: pair[1].w*pair[1].h)
                expected_bounds = pymupdf.Rect(shape.x*72-1, shape.y*72-1,
                                               (shape.x+shape.w)*72+1,
                                               (shape.y+shape.h)*72+1)
                if not expected_bounds.contains(bbox):
                    fit_errors.append((path.name, "Text clipped / overflow", span["text"],
                                       tuple(bbox), tuple(expected_bounds)))
                owners.append((owner, bbox, span["text"]))
            for i, (a, ra, ta) in enumerate(owners):
                for b, rb, tb in owners[i+1:]:
                    if a != b and ra.intersects(rb):
                        overlap = ra & rb
                        if overlap.width >= .5 and overlap.height >= .5:
                            fit_errors.append((path.name, "Text overlap", ta, tb, tuple(overlap)))
            results.append({"figure": item.name, "pdf": path.name, "pages": 1,
                            "vector_paths": vector_count,
                            "raster_images": 0, "visible_text_spans": len(spans),
                            "minimum_native_font_pt": round(smallest_font, 3),
                            "minimum_article_font_pt": round(article_font, 3)})
    assert not fit_errors, "\n" + "\n".join(str(error) for error in fit_errors)
    return results


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true",
                        help="Validate existing outputs without regenerating.")
    parser.add_argument("--preview", action="store_true",
                        help="Render inspection PNGs from the vector PDFs.")
    parser.add_argument("--no-export", action="store_true",
                        help="Generate only PPTX; skip PowerPoint COM and PDF checks.")
    args = parser.parse_args()
    items = figures()
    validate_model(items)
    if not args.check:
        write_deck(items)
        if not args.no_export:
            subprocess.run(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass",
                            "-File", str(ROOT / "export_powerpoint.ps1"),
                            "-DeckPath", str(DECK), "-OutputDirectory", str(ROOT)],
                           check=True)
    summary = {"pptx": DECK.name, "paper_title": PAPER_TITLE,
               "article_width_inches": PAPER_WIDTH, "slides": validate_deck(items),
               "export_route": "Microsoft PowerPoint 16 COM SaveAs PDF (native vectors)"}
    if not args.no_export:
        summary["pdf_exports"] = validate_pdfs(items, args.preview)
    summary["totals"] = {
        "slides": len(items),
        "native_shapes": sum(s["native_shapes"] for s in summary["slides"]),
        "native_connectors": sum(s["native_connectors"] for s in summary["slides"]),
        "native_text_runs": sum(s["native_text_runs"] for s in summary["slides"]),
        "pictures": 0,
    }
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
